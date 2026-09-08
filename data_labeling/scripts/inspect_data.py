#!/usr/bin/env python3
import csv
import json
from pathlib import Path
import pickle
import cv2

ROOT = Path(".").resolve()

VIDEO_EXTS = {".mp4", ".avi", ".mov", ".mkv", ".webm"}
ANN_EXTS = {".json", ".jsonl", ".csv", ".txt", ".pkl", ".pickle", ".npy", ".npz", ".xlsx"}

REPORTS = ROOT / "reports"
REPORTS.mkdir(exist_ok=True)


def safe_rel(path: Path) -> str:
    try:
        return str(path.resolve().relative_to(ROOT))
    except Exception:
        return str(path)


def video_info(path: Path):
    cap = cv2.VideoCapture(str(path))
    if not cap.isOpened():
        return {
            "readable": False,
            "frames": "",
            "fps": "",
            "width": "",
            "height": "",
            "duration_sec": "",
        }
    frames = int(cap.get(cv2.CAP_PROP_FRAME_COUNT))
    fps = float(cap.get(cv2.CAP_PROP_FPS) or 0)
    width = int(cap.get(cv2.CAP_PROP_FRAME_WIDTH))
    height = int(cap.get(cv2.CAP_PROP_FRAME_HEIGHT))
    cap.release()
    duration = frames / fps if fps > 0 else 0
    return {
        "readable": True,
        "frames": frames,
        "fps": fps,
        "width": width,
        "height": height,
        "duration_sec": duration,
    }


def guess_annotation_summary(path: Path):
    out = {
        "path": safe_rel(path),
        "suffix": path.suffix.lower(),
        "size_bytes": path.stat().st_size,
        "readable": False,
        "kind": "",
        "top_keys_or_columns": "",
        "num_items": "",
        "error": "",
    }

    try:
        suffix = path.suffix.lower()

        if suffix == ".json":
            obj = json.loads(path.read_text(errors="replace"))
            out["readable"] = True
            out["kind"] = type(obj).__name__
            if isinstance(obj, dict):
                out["top_keys_or_columns"] = ",".join(list(obj.keys())[:30])
                out["num_items"] = len(obj)
            elif isinstance(obj, list):
                out["num_items"] = len(obj)
                if obj and isinstance(obj[0], dict):
                    out["top_keys_or_columns"] = ",".join(list(obj[0].keys())[:30])

        elif suffix == ".jsonl":
            first = None
            n = 0
            with path.open("r", errors="replace") as f:
                for line in f:
                    if line.strip():
                        n += 1
                        if first is None:
                            first = json.loads(line)
            out["readable"] = True
            out["kind"] = "jsonl"
            out["num_items"] = n
            if isinstance(first, dict):
                out["top_keys_or_columns"] = ",".join(list(first.keys())[:30])

        elif suffix == ".csv":
            with path.open("r", errors="replace", newline="") as f:
                reader = csv.reader(f)
                header = next(reader, [])
                n = sum(1 for _ in reader)
            out["readable"] = True
            out["kind"] = "csv"
            out["top_keys_or_columns"] = ",".join(header[:30])
            out["num_items"] = n

        elif suffix == ".xlsx":
            from openpyxl import load_workbook

            workbook = load_workbook(path, read_only=True, data_only=True)
            sheet_names = workbook.sheetnames
            out["readable"] = True
            out["kind"] = "xlsx"
            out["top_keys_or_columns"] = ",".join(sheet_names[:30])
            out["num_items"] = sum(workbook[name].max_row for name in sheet_names)
            workbook.close()

        elif suffix in {".pkl", ".pickle"}:
            with path.open("rb") as f:
                obj = pickle.load(f)
            out["readable"] = True
            out["kind"] = type(obj).__name__
            if isinstance(obj, dict):
                out["top_keys_or_columns"] = ",".join([str(k) for k in list(obj.keys())[:30]])
                out["num_items"] = len(obj)
            elif isinstance(obj, (list, tuple)):
                out["num_items"] = len(obj)

        elif suffix == ".txt":
            lines = path.read_text(errors="replace").splitlines()
            out["readable"] = True
            out["kind"] = "txt"
            out["num_items"] = len(lines)
            out["top_keys_or_columns"] = lines[0][:200] if lines else ""

        else:
            out["kind"] = suffix.lstrip(".")
            out["readable"] = True

    except Exception as e:
        out["error"] = repr(e)

    return out


def main():
    videos = []
    anns = []

    for path in ROOT.rglob("*"):
        if not path.is_file():
            continue

        if ".git" in path.parts or ".venv" in path.parts:
            continue

        suffix = path.suffix.lower()

        if suffix in VIDEO_EXTS:
            info = video_info(path)
            info.update({
                "path": safe_rel(path),
                "size_bytes": path.stat().st_size,
            })
            videos.append(info)

        elif suffix in ANN_EXTS:
            anns.append(guess_annotation_summary(path))

    videos.sort(key=lambda x: x["path"])
    anns.sort(key=lambda x: x["path"])

    with (REPORTS / "video_inventory.csv").open("w", newline="") as f:
        fieldnames = ["path", "size_bytes", "readable", "frames", "fps", "width", "height", "duration_sec"]
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(videos)

    with (REPORTS / "annotation_inventory.csv").open("w", newline="") as f:
        fieldnames = ["path", "suffix", "size_bytes", "readable", "kind", "top_keys_or_columns", "num_items", "error"]
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(anns)

    summary = {
        "num_videos": len(videos),
        "num_readable_videos": sum(1 for v in videos if v["readable"]),
        "num_annotation_like_files": len(anns),
        "num_readable_annotation_like_files": sum(1 for a in anns if a["readable"]),
        "videos_by_root": {},
        "annotations_by_root": {},
    }

    for v in videos:
        root = v["path"].split("/")[0]
        summary["videos_by_root"][root] = summary["videos_by_root"].get(root, 0) + 1

    for a in anns:
        root = a["path"].split("/")[0]
        summary["annotations_by_root"][root] = summary["annotations_by_root"].get(root, 0) + 1

    (REPORTS / "inspection_summary.json").write_text(json.dumps(summary, indent=2), encoding="utf-8")

    print(json.dumps(summary, indent=2))


if __name__ == "__main__":
    main()
